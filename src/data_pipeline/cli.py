"""Commandes du pipeline Bottleneck exécutées localement ou par Kestra."""

from __future__ import annotations

import argparse
import json
import math
import os
import re
import sys
from contextlib import contextmanager
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable
from decimal import Decimal, ROUND_HALF_UP

# Le dossier de données ./bottleneck masque sinon le package optionnel homonyme
# que Pandas tente de détecter lors de son import depuis la racine du dépôt.
_repository_root = Path(__file__).resolve().parents[2]
_original_sys_path = sys.path[:]
sys.path = [
    item for item in sys.path if Path(item or ".").resolve() != _repository_root
]
try:
    import pandas as pd
finally:
    sys.path = _original_sys_path

import duckdb

PROJECT_ROOT = _repository_root
DEFAULT_SQL_DIR = PROJECT_ROOT / "sql"

SOURCE_FILES = {
    "erp": "Fichier_erp.xlsx",
    "liaison": "fichier_liaison.xlsx",
    "web": "Fichier_web.xlsx",
}

REQUIRED_COLUMNS = {
    "erp": {"product_id", "onsale_web", "price", "stock_quantity", "stock_status"},
    "liaison": {"product_id", "id_web"},
    "web": {"sku", "total_sales", "post_title", "post_type", "post_modified", "guid"},
}

EXPECTED_ROW_COUNTS = {
    "erp_after_deduplication": 825,
    "liaison_after_deduplication": 825,
    "web_after_cleaning": 1428,
    "web_after_deduplication": 714,
}


def normalize_identifier(value: object) -> str | None:
    """Normalise une clé Excel sans perdre les identifiants alphanumériques."""

    if pd.isna(value):
        return None
    if isinstance(value, bool):
        return str(value).lower()
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))

    text = str(value).strip()
    if not text:
        return None
    if re.fullmatch(r"-?\d+\.0", text):
        return text[:-2]
    return text


def _check_columns(name: str, dataframe: pd.DataFrame) -> None:
    missing = REQUIRED_COLUMNS[name] - set(dataframe.columns)
    if missing:
        raise ValueError(
            f"Colonnes absentes dans {SOURCE_FILES[name]} : {', '.join(sorted(missing))}"
        )


def stage_excel_sources(
    input_dir: Path, staging_dir: Path
) -> dict[str, dict[str, int]]:
    """Convertit les classeurs Excel en CSV bruts lisibles par DuckDB."""

    input_dir = input_dir.resolve()
    staging_dir.mkdir(parents=True, exist_ok=True)
    manifest: dict[str, dict[str, int]] = {}

    for source_name, file_name in SOURCE_FILES.items():
        file_path = input_dir / file_name
        if not file_path.exists():
            raise FileNotFoundError(f"Fichier source introuvable : {file_path}")

        dataframe = pd.read_excel(file_path, engine="openpyxl", dtype=object)
        dataframe.columns = [str(column).strip() for column in dataframe.columns]
        _check_columns(source_name, dataframe)

        for identifier in ("product_id", "id_web", "sku"):
            if identifier in dataframe.columns:
                dataframe[identifier] = dataframe[identifier].map(normalize_identifier)

        output_path = staging_dir / f"{source_name}.csv"
        dataframe.to_csv(
            output_path, index=False, encoding="utf-8", lineterminator="\n"
        )
        manifest[source_name] = {
            "rows": int(len(dataframe)),
            "columns": int(len(dataframe.columns)),
            "full_row_duplicates": int(dataframe.duplicated().sum()),
        }

    (staging_dir / "staging_manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return manifest


@contextmanager
def working_directory(path: Path) -> Iterable[None]:
    previous = Path.cwd()
    os.chdir(path)
    try:
        yield
    finally:
        os.chdir(previous)


def resolve_from_work_dir(path: Path, work_dir: Path) -> Path:
    return path if path.is_absolute() else (work_dir / path).resolve()


def execute_sql_script(database: Path, script: Path, work_dir: Path) -> None:
    """Exécute un fichier SQL DuckDB dans le répertoire de travail partagé."""

    work_dir = work_dir.resolve()
    database = resolve_from_work_dir(database, work_dir)
    script = script.resolve()
    database.parent.mkdir(parents=True, exist_ok=True)
    if not script.exists():
        raise FileNotFoundError(f"Script SQL introuvable : {script}")

    with working_directory(work_dir):
        with duckdb.connect(str(database)) as connection:
            connection.execute(script.read_text(encoding="utf-8"))


def segment_wines(
    database: Path, output_dir: Path, work_dir: Path
) -> dict[str, float | int]:
    """Calcule le z-score et extrait vins premium et vins ordinaires."""

    work_dir = work_dir.resolve()
    database = resolve_from_work_dir(database, work_dir)
    output_dir = resolve_from_work_dir(output_dir, work_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    with duckdb.connect(str(database)) as connection:
        wines = connection.execute("""
            SELECT
                revenue.product_id,
                revenue.id_web,
                revenue.post_title,
                revenue.price,
                revenue.total_sales,
                revenue.revenue
            FROM revenue_by_product AS revenue
            ORDER BY revenue.product_id
            """).fetch_df()

        if wines.empty:
            raise ValueError(
                "La table fusionnée est vide : impossible de calculer les z-scores."
            )

        mean_price = float(wines["price"].mean())
        std_price = float(wines["price"].std(ddof=0))
        if math.isclose(std_price, 0.0):
            raise ValueError(
                "L'écart type des prix est nul : le z-score n'est pas défini."
            )

        wines["z_score"] = (wines["price"] - mean_price) / std_price
        wines["category"] = (
            wines["z_score"].gt(2).map({True: "premium", False: "ordinaire"})
        )
        wines["wine_type"] = (
            wines["z_score"].gt(2).map({True: "millésime", False: "ordinaire"})
        )

        connection.register("segmentation_dataframe", wines)
        connection.execute(
            "CREATE OR REPLACE TABLE wine_segmentation AS "
            "SELECT * FROM segmentation_dataframe"
        )
        connection.unregister("segmentation_dataframe")

    export_columns = [
        "product_id",
        "id_web",
        "post_title",
        "price",
        "total_sales",
        "revenue",
        "z_score",
        "category",
        "wine_type",
    ]
    premium = wines.loc[wines["category"] == "premium", export_columns].copy()
    ordinary = wines.loc[wines["category"] == "ordinaire", export_columns].copy()
    for dataframe in (premium, ordinary):
        dataframe["price"] = dataframe["price"].round(2)
        dataframe["revenue"] = dataframe["revenue"].round(2)
        dataframe["z_score"] = dataframe["z_score"].round(6)

    premium.sort_values("z_score", ascending=False).to_csv(
        output_dir / "vins_premium.csv", index=False, encoding="utf-8-sig"
    )
    ordinary.sort_values("product_id").to_csv(
        output_dir / "vins_ordinaires.csv", index=False, encoding="utf-8-sig"
    )

    metrics: dict[str, float | int] = {
        "mean_price": round(mean_price, 6),
        "population_standard_deviation": round(std_price, 6),
        "premium_count": int(len(premium)),
        "ordinary_count": int(len(ordinary)),
    }
    (output_dir / "zscore_metrics.json").write_text(
        json.dumps(metrics, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return metrics


def _write_legacy_xls(
    report_path: Path, revenue: pd.DataFrame, summary: dict[str, float | int]
) -> None:
    import xlwt

    workbook = xlwt.Workbook(encoding="utf-8")
    summary_sheet = workbook.add_sheet("Synthèse")
    detail_sheet = workbook.add_sheet("CA par produit")

    title_style = xlwt.easyxf(
        "font: bold on, height 320, colour white;"
        "pattern: pattern solid, fore_colour dark_blue;"
        "align: horiz center, vert center;"
    )
    label_style = xlwt.easyxf(
        "font: bold on, colour white; pattern: pattern solid, fore_colour blue_gray;"
    )
    header_style = xlwt.easyxf(
        "font: bold on, colour white; pattern: pattern solid, fore_colour dark_blue;"
        "align: horiz center; borders: bottom thin, bottom_colour gray40;"
    )
    currency_style = xlwt.easyxf(num_format_str="€ #,##0.00")
    integer_style = xlwt.easyxf(num_format_str="#,##0")

    summary_sheet.write_merge(0, 0, 0, 5, "Rapport du chiffre d'affaires", title_style)
    summary_rows = [
        ("Chiffre d'affaires total", float(summary["total_revenue"]), currency_style),
        ("Nombre de produits", int(summary["product_count"]), integer_style),
        ("Unités vendues", int(summary["bottles_sold"]), integer_style),
    ]
    for row_index, (label, value, style) in enumerate(summary_rows, start=2):
        summary_sheet.write(row_index, 0, label, label_style)
        summary_sheet.write(row_index, 1, value, style)
    summary_sheet.write(7, 0, "Source")
    summary_sheet.write(7, 1, "Fusion ERP → liaison → web ; CA = prix × ventes")
    summary_sheet.col(0).width = 9000
    summary_sheet.col(1).width = 9000

    headers = [
        "product_id",
        "id_web",
        "Produit",
        "Prix unitaire (€)",
        "Quantité vendue",
        "Chiffre d'affaires (€)",
    ]
    for column_index, header in enumerate(headers):
        detail_sheet.write(0, column_index, header, header_style)

    for row_index, row in enumerate(revenue.itertuples(index=False), start=1):
        detail_sheet.write(row_index, 0, int(row.product_id))
        detail_sheet.write(row_index, 1, str(row.id_web))
        detail_sheet.write(row_index, 2, str(row.post_title))
        detail_sheet.write(row_index, 3, float(row.price), currency_style)
        detail_sheet.write(row_index, 4, int(row.total_sales), integer_style)
        detail_sheet.write(row_index, 5, float(row.revenue), currency_style)

    for column, width in enumerate([3500, 4500, 18000, 5000, 5000, 6500]):
        detail_sheet.col(column).width = width
    detail_sheet.set_panes_frozen(True)
    detail_sheet.set_horz_split_pos(1)
    detail_sheet.set_remove_splits(True)
    workbook.save(str(report_path))


def _write_modern_xlsx(
    report_path: Path, revenue: pd.DataFrame, summary: dict[str, float | int]
) -> None:
    from openpyxl import Workbook
    from openpyxl.formatting.rule import DataBarRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.table import Table, TableStyleInfo

    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Synthèse"
    detail_sheet = workbook.create_sheet("CA par produit")

    dark_blue = "17324D"
    medium_blue = "365B73"
    gold = "D6A84B"
    light_blue = "EAF1F5"
    white = "FFFFFF"
    thin_gray = Side(style="thin", color="D7E0E5")

    summary_sheet.sheet_view.showGridLines = False
    summary_sheet.merge_cells("A1:F1")
    summary_sheet["A1"] = "Rapport du chiffre d'affaires"
    summary_sheet["A1"].font = Font(
        name="Aptos Display", size=18, bold=True, color=white
    )
    summary_sheet["A1"].fill = PatternFill("solid", fgColor=dark_blue)
    summary_sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    summary_sheet.row_dimensions[1].height = 32

    cards = [
        (
            "A3",
            "Chiffre d'affaires total",
            float(summary["total_revenue"]),
            "€ #,##0.00",
        ),
        ("C3", "Nombre de produits", int(summary["product_count"]), "#,##0"),
        ("E3", "Unités vendues", int(summary["bottles_sold"]), "#,##0"),
    ]
    for anchor, label, value, number_format in cards:
        cell = summary_sheet[anchor]
        cell.value = label
        cell.font = Font(name="Aptos", size=10, bold=True, color=medium_blue)
        value_cell = summary_sheet.cell(row=cell.row + 1, column=cell.column)
        value_cell.value = value
        value_cell.number_format = number_format
        value_cell.font = Font(
            name="Aptos Display", size=16, bold=True, color=dark_blue
        )
        for row in (cell.row, cell.row + 1):
            for column in (cell.column, cell.column + 1):
                target = summary_sheet.cell(row=row, column=column)
                target.fill = PatternFill("solid", fgColor=light_blue)
                target.border = Border(bottom=thin_gray)
        summary_sheet.merge_cells(
            start_row=cell.row,
            start_column=cell.column,
            end_row=cell.row,
            end_column=cell.column + 1,
        )
        summary_sheet.merge_cells(
            start_row=cell.row + 1,
            start_column=cell.column,
            end_row=cell.row + 1,
            end_column=cell.column + 1,
        )

    summary_sheet["A7"] = "Méthode"
    summary_sheet["A7"].font = Font(bold=True, color=dark_blue)
    summary_sheet["A8"] = (
        "Fusion ERP → liaison → web ; chiffre d'affaires = prix ERP × ventes web."
    )
    summary_sheet.merge_cells("A8:F8")
    summary_sheet["A8"].alignment = Alignment(wrap_text=True)
    for column in range(1, 7):
        summary_sheet.column_dimensions[chr(64 + column)].width = 16

    detail_sheet.sheet_view.showGridLines = False
    detail_sheet.merge_cells("A1:F1")
    detail_sheet["A1"] = "Chiffre d'affaires par produit"
    detail_sheet["A1"].font = Font(
        name="Aptos Display", size=16, bold=True, color=white
    )
    detail_sheet["A1"].fill = PatternFill("solid", fgColor=dark_blue)
    detail_sheet["A1"].alignment = Alignment(horizontal="center")
    detail_sheet.row_dimensions[1].height = 28

    headers = [
        "product_id",
        "id_web",
        "Produit",
        "Prix unitaire (€)",
        "Quantité vendue",
        "Chiffre d'affaires (€)",
    ]
    detail_sheet.append([])
    detail_sheet.append(headers)
    for row in revenue.itertuples(index=False):
        detail_sheet.append(
            [
                int(row.product_id),
                str(row.id_web),
                str(row.post_title),
                float(row.price),
                int(row.total_sales),
                float(row.revenue),
            ]
        )

    last_row = len(revenue) + 3
    table = Table(displayName="RevenueByProduct", ref=f"A3:F{last_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    detail_sheet.add_table(table)
    detail_sheet.freeze_panes = "A4"
    detail_sheet.column_dimensions["A"].width = 14
    detail_sheet.column_dimensions["B"].width = 16
    detail_sheet.column_dimensions["C"].width = 72
    detail_sheet.column_dimensions["D"].width = 19
    detail_sheet.column_dimensions["E"].width = 19
    detail_sheet.column_dimensions["F"].width = 24
    for row in range(4, last_row + 1):
        detail_sheet.cell(row, 4).number_format = "€ #,##0.00"
        detail_sheet.cell(row, 5).number_format = "#,##0"
        detail_sheet.cell(row, 6).number_format = "€ #,##0.00"
    detail_sheet.conditional_formatting.add(
        f"F4:F{last_row}",
        DataBarRule(start_type="min", end_type="max", color=gold, showValue=True),
    )

    workbook.save(report_path)


def export_revenue_reports(
    database: Path, output_dir: Path, work_dir: Path
) -> dict[str, float | int]:
    """Exporte le CA par produit et le CA global en formats Excel moderne et legacy."""

    work_dir = work_dir.resolve()
    database = resolve_from_work_dir(database, work_dir)
    output_dir = resolve_from_work_dir(output_dir, work_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    with duckdb.connect(str(database), read_only=True) as connection:
        revenue = connection.execute(
            "SELECT * FROM revenue_by_product ORDER BY revenue DESC, product_id"
        ).fetch_df()
        summary_row = connection.execute("SELECT * FROM revenue_summary").fetchone()

    if summary_row is None:
        raise ValueError("La synthèse du chiffre d'affaires est absente.")
    summary: dict[str, float | int] = {
        "product_count": int(summary_row[0]),
        "bottles_sold": int(summary_row[1]),
        "total_revenue": float(summary_row[2]),
    }
    _write_legacy_xls(output_dir / "rapport_chiffre_affaires.xls", revenue, summary)
    _write_modern_xlsx(output_dir / "rapport_chiffre_affaires.xlsx", revenue, summary)
    return summary


VALIDATION_STAGES = ("cleaning", "merge", "revenue", "segmentation")

EXPECTED_MERGED_COUNT = 714
EXPECTED_TOTAL_REVENUE = Decimal("70568.60")
EXPECTED_PREMIUM_COUNT = 30


def _money(value: object) -> Decimal:
    """Compare les montants au centime sans masquer un écart d'un centime."""
    amount = Decimal(str(value))
    if not amount.is_finite():
        raise ValueError(f"Montant non fini : {value}")
    return amount.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def validate_pipeline(
    database: Path,
    output_dir: Path,
    work_dir: Path,
    stage: str = "all",
) -> dict[str, object]:
    """Exécute les contrôles bloquants d'une étape ou du pipeline entier."""

    if stage not in (*VALIDATION_STAGES, "all"):
        raise ValueError(f"Étape de validation inconnue : {stage}")

    work_dir = work_dir.resolve()
    database = resolve_from_work_dir(database, work_dir)
    output_dir = resolve_from_work_dir(output_dir, work_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    stages = VALIDATION_STAGES if stage == "all" else (stage,)
    checks = {}
    metrics = {}
    row_count_tests = {}

    with duckdb.connect(str(database), read_only=True) as connection:

        def scalar(query):
            return connection.execute(query).fetchone()[0]

        def check_zero(name, query):
            actual = int(scalar(query))
            metrics[name] = actual
            checks[name] = actual == 0

        def check_count(name, table, expected):
            actual = int(scalar(f"SELECT COUNT(*) FROM {table}"))
            metrics[name] = actual
            row_count_tests[name] = {
                "actual": actual,
                "expected": expected,
                "passed": actual == expected,
            }
            checks[f"{name}_matches_expected"] = actual == expected

        if "cleaning" in stages:
            tables = {
                "erp_after_deduplication": "erp_clean",
                "liaison_after_deduplication": "liaison_clean",
                "web_after_cleaning": "web_not_null",
                "web_after_deduplication": "web_clean",
            }
            for name, table in tables.items():
                check_count(name, table, EXPECTED_ROW_COUNTS[name])

            # Les NULL de liaison.id_web sont autorisés.
            required_fields = {
                "erp_not_null": """
                    product_id IS NULL OR price IS NULL
                    OR NOT isfinite(price)
                """,
                "erp_clean": """
                    product_id IS NULL OR price IS NULL
                    OR NOT isfinite(price)
                """,
                "liaison_valid": "product_id IS NULL",
                "liaison_clean": "product_id IS NULL",
                "web_not_null": """
                    NULLIF(TRIM(sku), '') IS NULL
                    OR total_sales IS NULL
                    OR NULLIF(TRIM(post_title), '') IS NULL
                    OR post_type IS NULL
                    OR post_type NOT IN ('product', 'attachment')
                """,
                "web_clean": """
                    NULLIF(TRIM(sku), '') IS NULL
                    OR total_sales IS NULL
                    OR NULLIF(TRIM(post_title), '') IS NULL
                    OR post_type IS NULL
                    OR post_type NOT IN ('product', 'attachment')
                """,
            }
            for table, condition in required_fields.items():
                check_zero(
                    f"{table}_invalid_required_fields",
                    f"SELECT COUNT(*) FROM {table} WHERE {condition}",
                )

            for table, column in (
                ("erp_clean", "product_id"),
                ("liaison_clean", "product_id"),
                ("liaison_clean", "id_web"),
                ("web_clean", "sku"),
            ):
                check_zero(
                    f"{table}_{column}_duplicates",
                    f"""
                    SELECT COUNT(*)
                    FROM (
                        SELECT {column}
                        FROM {table}
                        WHERE {column} IS NOT NULL
                        GROUP BY {column}
                        HAVING COUNT(*) > 1
                    ) AS duplicates
                    """,
                )

        if "merge" in stages:
            check_count("merged_count", "merged_wines", EXPECTED_MERGED_COUNT)

            for column in ("product_id", "id_web"):
                check_zero(
                    f"merged_{column}_duplicates",
                    f"""
                    SELECT COUNT(*)
                    FROM (
                        SELECT {column}
                        FROM merged_wines
                        WHERE {column} IS NOT NULL
                        GROUP BY {column}
                        HAVING COUNT(*) > 1
                    ) AS duplicates
                    """,
                )

            check_zero(
                "merged_invalid_required_fields",
                """
                SELECT COUNT(*)
                FROM merged_wines
                WHERE product_id IS NULL
                   OR NULLIF(TRIM(id_web), '') IS NULL
                   OR NULLIF(TRIM(sku), '') IS NULL
                   OR NULLIF(TRIM(post_title), '') IS NULL
                   OR price IS NULL
                   OR NOT isfinite(price)
                   OR total_sales IS NULL
                   OR id_web <> sku
                """,
            )

        if "revenue" in stages:
            check_count(
                "revenue_product_count",
                "revenue_by_product",
                EXPECTED_MERGED_COUNT,
            )

            check_zero(
                "revenue_product_id_duplicates",
                """
                SELECT COUNT(*)
                FROM (
                    SELECT product_id
                    FROM revenue_by_product
                    GROUP BY product_id
                    HAVING COUNT(*) > 1
                ) AS duplicates
                """,
            )

            # Comparaison avec la jointure, pas seulement avec la synthèse.
            check_zero(
                "revenue_detail_mismatches",
                """
                SELECT COUNT(*)
                FROM merged_wines AS m
                FULL OUTER JOIN revenue_by_product AS r
                    ON m.product_id = r.product_id
                WHERE m.product_id IS NULL
                   OR r.product_id IS NULL
                   OR r.id_web IS DISTINCT FROM m.id_web
                   OR r.price IS DISTINCT FROM m.price
                   OR r.total_sales IS DISTINCT FROM m.total_sales
                   OR r.revenue IS NULL
                   OR NOT isfinite(r.revenue)
                   OR r.revenue IS DISTINCT FROM
                      ROUND(m.price * m.total_sales, 2)
                """,
            )

            total_detail = scalar("SELECT SUM(revenue) FROM revenue_by_product")
            summaries = connection.execute("""
                SELECT product_count, bottles_sold, total_revenue
                FROM revenue_summary
                """).fetchall()

            checks["revenue_summary_has_one_row"] = len(summaries) == 1
            total_summary = summaries[0][2] if len(summaries) == 1 else None

            def matches_money(left, right):
                try:
                    return _money(left) == _money(right)
                except (ValueError, ArithmeticError):
                    return False

            checks["revenue_reconciles"] = matches_money(total_detail, total_summary)
            checks["revenue_detail_matches_reference"] = matches_money(
                total_detail, EXPECTED_TOTAL_REVENUE
            )
            checks["revenue_summary_matches_reference"] = matches_money(
                total_summary, EXPECTED_TOTAL_REVENUE
            )

            expected_summary = connection.execute(
                "SELECT COUNT(*), SUM(total_sales) FROM merged_wines"
            ).fetchone()
            checks["revenue_summary_counts_match"] = (
                len(summaries) == 1 and summaries[0][:2] == expected_summary
            )

            # Chaînes pour conserver exactement la représentation monétaire.
            metrics["total_revenue"] = (
                str(total_summary) if total_summary is not None else None
            )
            metrics["detail_total_revenue"] = (
                str(total_detail) if total_detail is not None else None
            )
            metrics["expected_total_revenue"] = str(EXPECTED_TOTAL_REVENUE)

        if "segmentation" in stages:
            check_count(
                "segmentation_count",
                "wine_segmentation",
                EXPECTED_MERGED_COUNT,
            )

            premium_count = int(scalar("""
                SELECT COUNT(*) FROM wine_segmentation
                WHERE category = 'premium'
                """))
            ordinary_count = int(scalar("""
                SELECT COUNT(*) FROM wine_segmentation
                WHERE category = 'ordinaire'
                """))
            metrics["premium_count"] = premium_count
            metrics["ordinary_count"] = ordinary_count

            checks["premium_count_matches_reference"] = (
                premium_count == EXPECTED_PREMIUM_COUNT
            )
            checks["ordinary_count_matches_reference"] = (
                ordinary_count == EXPECTED_MERGED_COUNT - EXPECTED_PREMIUM_COUNT
            )

            check_zero(
                "segmentation_product_id_duplicates",
                """
                SELECT COUNT(*)
                FROM (
                    SELECT product_id
                    FROM wine_segmentation
                    GROUP BY product_id
                    HAVING COUNT(*) > 1
                ) AS duplicates
                """,
            )

            check_zero(
                "segmentation_invalid_classification",
                """
                SELECT COUNT(*)
                FROM wine_segmentation
                WHERE z_score IS NULL
                   OR NOT isfinite(z_score)
                   OR category IS NULL
                   OR category NOT IN ('premium', 'ordinaire')
                   OR (category = 'premium' AND z_score <= 2)
                   OR (category = 'ordinaire' AND z_score > 2)
                """,
            )

            # Vérifie également le calcul du z-score indépendamment du
            # traitement Pandas, avec le même écart type population (ddof=0).
            stats = connection.execute("""
                SELECT AVG(price), STDDEV_POP(price)
                FROM revenue_by_product
                """).fetchone()

            stats_valid = (
                stats[0] is not None
                and stats[1] is not None
                and math.isfinite(float(stats[0]))
                and math.isfinite(float(stats[1]))
                and float(stats[1]) > 0
            )
            checks["zscore_statistics_valid"] = stats_valid

            check_zero(
                "segmentation_source_mismatches",
                """
                SELECT COUNT(*)
                FROM revenue_by_product AS r
                FULL OUTER JOIN wine_segmentation AS s
                    ON r.product_id = s.product_id
                WHERE r.product_id IS NULL
                   OR s.product_id IS NULL
                   OR s.price IS DISTINCT FROM r.price
                   OR s.id_web IS DISTINCT FROM r.id_web
                """,
            )

            if stats_valid:
                invalid_zscores = int(
                    connection.execute(
                        """
                    SELECT COUNT(*)
                    FROM wine_segmentation AS s
                    JOIN revenue_by_product AS r
                        ON s.product_id = r.product_id
                    WHERE s.z_score IS NULL
                       OR NOT isfinite(s.z_score)
                       OR ABS(s.z_score - ((r.price - ?) / ?)) > 1e-9
                    """,
                        [float(stats[0]), float(stats[1])],
                    ).fetchone()[0]
                )
                metrics["incorrect_zscores"] = invalid_zscores
                checks["zscore_calculation_is_correct"] = invalid_zscores == 0
            else:
                checks["zscore_calculation_is_correct"] = False

    report = {
        "status": "passed" if all(checks.values()) else "failed",
        "stage": stage,
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "metrics": metrics,
        "row_count_tests": row_count_tests,
        "checks": checks,
    }

    filename = "quality_report.json" if stage == "all" else f"quality_{stage}.json"
    (output_dir / filename).write_text(
        json.dumps(report, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    if report["status"] != "passed":
        failed = [name for name, passed in checks.items() if not passed]
        message = f"Contrôles qualité en échec [{stage}] : {', '.join(failed)}"
        print(message, file=sys.stderr, flush=True)
        raise ValueError(message)

    return report


def run_pipeline(
    input_dir: Path,
    work_dir: Path,
    sql_dir: Path = DEFAULT_SQL_DIR,
) -> None:
    """Exécute les traitements avec validation bloquante par étape."""

    work_dir = work_dir.resolve()
    sql_dir = sql_dir.resolve()
    work_dir.mkdir(parents=True, exist_ok=True)

    stage_excel_sources(input_dir, work_dir / "staging")
    database = Path("pipeline.duckdb")
    output_dir = Path("outputs")

    steps = (
        ("00_load_raw.sql", None),
        ("01_remove_missing_values.sql", None),
        ("02_deduplicate.sql", "cleaning"),
        ("03_merge_systems.sql", "merge"),
        ("04_calculate_revenue.sql", "revenue"),
    )

    for script_name, validation_stage in steps:
        execute_sql_script(database, sql_dir / script_name, work_dir)
        if validation_stage:
            validate_pipeline(database, output_dir, work_dir, stage=validation_stage)

    segment_wines(database, output_dir, work_dir)
    validate_pipeline(database, output_dir, work_dir, stage="segmentation")
    validate_pipeline(database, output_dir, work_dir)
    export_revenue_reports(database, output_dir, work_dir)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    subparsers = parser.add_subparsers(dest="command", required=True)

    stage = subparsers.add_parser(
        "stage", help="Convertir les sources Excel en CSV bruts"
    )
    stage.add_argument("--input-dir", type=Path, required=True)
    stage.add_argument("--staging-dir", type=Path, default=Path("staging"))

    sql = subparsers.add_parser("sql", help="Exécuter un script SQL DuckDB")
    sql.add_argument("--database", type=Path, default=Path("pipeline.duckdb"))
    sql.add_argument("--script", type=Path, required=True)
    sql.add_argument("--work-dir", type=Path, default=Path("."))

    segment = subparsers.add_parser(
        "segment", help="Calculer les z-scores et exporter les segments"
    )
    segment.add_argument("--database", type=Path, default=Path("pipeline.duckdb"))
    segment.add_argument("--output-dir", type=Path, default=Path("outputs"))
    segment.add_argument("--work-dir", type=Path, default=Path("."))

    export = subparsers.add_parser(
        "export", help="Exporter les rapports de chiffre d'affaires"
    )
    export.add_argument("--database", type=Path, default=Path("pipeline.duckdb"))
    export.add_argument("--output-dir", type=Path, default=Path("outputs"))
    export.add_argument("--work-dir", type=Path, default=Path("."))

    validate = subparsers.add_parser("validate", help="Exécuter les contrôles qualité")
    validate.add_argument("--database", type=Path, default=Path("pipeline.duckdb"))
    validate.add_argument("--output-dir", type=Path, default=Path("outputs"))
    validate.add_argument("--work-dir", type=Path, default=Path("."))
    validate.add_argument(
        "--stage",
        choices=["all", *VALIDATION_STAGES],
        default="all",
    )

    run = subparsers.add_parser("run", help="Exécuter le pipeline complet sans Kestra")
    run.add_argument("--input-dir", type=Path, required=True)
    run.add_argument("--work-dir", type=Path, default=Path("build"))
    run.add_argument("--sql-dir", type=Path, default=DEFAULT_SQL_DIR)
    return parser


def main() -> None:
    args = build_parser().parse_args()
    if args.command == "stage":
        print(
            json.dumps(stage_excel_sources(args.input_dir, args.staging_dir), indent=2)
        )
    elif args.command == "sql":
        execute_sql_script(args.database, args.script, args.work_dir)
        print(f"Script exécuté : {args.script}")
    elif args.command == "segment":
        print(
            json.dumps(
                segment_wines(args.database, args.output_dir, args.work_dir), indent=2
            )
        )
    elif args.command == "export":
        print(
            json.dumps(
                export_revenue_reports(args.database, args.output_dir, args.work_dir),
                indent=2,
            )
        )
    elif args.command == "validate":
        print(
            json.dumps(
                validate_pipeline(
                    args.database,
                    args.output_dir,
                    args.work_dir,
                    stage=args.stage,
                ),
                indent=2,
            )
        )
    elif args.command == "run":
        run_pipeline(args.input_dir, args.work_dir, args.sql_dir)
        print(f"Pipeline terminé. Sorties : {(args.work_dir / 'outputs').resolve()}")


if __name__ == "__main__":
    main()
